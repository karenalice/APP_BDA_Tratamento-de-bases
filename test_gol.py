import unittest
import zipfile
import io
import re
from pathlib import Path
from tempfile import TemporaryDirectory

import pandas as pd

from gol_processor import (format_gol, merge_period, prepare_complement, recover_emission_dates,
                           report_rows, report_matrix, report_layouts, period_titles,
                           adjust_shares, write_zip, default_filters, filter_options,
                           daily_comparison_tables, daily_comparison_dataframe)


class GolFlowTests(unittest.TestCase):
    def sample(self):
        return pd.DataFrame({
            "DATA_EMISSAO": ["05/10/2026"] * 5,
            "DATA_EMBARQUE": ["20/10/2026"] * 5,
            "CIA_BILHETE": ["JJ", "G3", "G3", "JJ", "G3"],
            "CUPOM": ["1"] * 5,
            "BILHETE": [f"1234567890{i:03}" for i in range(5)],
            "TIPO": ["N"] * 5,
            "ORIGEM": ["CGB"] * 5,
            "DESTINO": ["GRU"] * 5,
            "TRECHO": ["CGB/GRU"] * 5,
            "IATA_AGENCIA": ["57500255", "57501255", "57500003", "57503255", "57504255"],
            "VALOR_BRL": ["100,00"] * 5,
            "BASE_VENDA": ["CGB", "CGB", "ART", "ART", "CGB"],
            "NUM_VOO": ["1234"] * 5,
            "CABINE": ["Y"] * 5,
            "CONSOLIDADA": ["1"] * 5,
            "CONSOLIDADA.1": ["ANTIGA"] * 5,
            "GOV": ["Nao", "Nao", "Nao", "Nao", "Sim"],
        })

    def test_format_and_report_rules(self):
        formatted, excluded, _, stats = format_gol(self.sample(), {57501255})
        self.assertEqual(len(excluded), 0)
        self.assertEqual(stats["cgb_rows"], 1)
        self.assertEqual(stats["gol_artmax_rows"], 1)
        self.assertEqual(stats["gol_artmax_iata_255_rows"], 0)
        self.assertEqual(str(formatted["IATA_AGENCIA"].dtype), "Int64")
        self.assertEqual(formatted.loc[2, "BASE_VENDA"], "CGB")
        self.assertEqual(formatted.loc[2, "CONSOLIDADA_1"], "43")
        self.assertEqual(formatted.loc[2, "CONSOLIDADA_2"], "ARARAUNA TURISMO")
        self.assertEqual(formatted.loc[3, "BASE_VENDA"], "ART")
        self.assertEqual(len(report_rows(formatted, "CONS")), 2)
        self.assertEqual(len(report_rows(formatted, "CORP")), 1)
        self.assertEqual(len(report_rows(formatted, "GOV")), 1)

    def test_adjust_only_original_jj_ad_and_zip(self):
        formatted, excluded, coupons, stats = format_gol(self.sample(), {57501255})
        extra = formatted.iloc[[0]].copy()
        extra["BILHETE"] = "9999999999999"
        extra["DATA_EMISSAO"] = pd.Timestamp("2026-09-01")
        full = merge_period(formatted.assign(_SOURCE_PERIOD="ORIGINAL"), extra)
        week = {"2026-10-05"}
        before = report_matrix(full, week, {})
        target = min(1, before[1]["shares"]["G3"] + 0.05)
        final, removed, prior, after = adjust_shares(full, week,
            {"CONS": {"general": target, "week": None, "advp": {}}}, set(), {})
        self.assertFalse(removed["CIA_BILHETE"].astype(str).eq("G3").any())
        self.assertFalse(removed.get("_SOURCE_PERIOD", pd.Series(dtype=str)).eq("COMPLEMENTAR").any())
        self.assertEqual(len(final) + len(removed), len(full))
        with TemporaryDirectory() as temp:
            original = Path(temp) / "original.csv"
            self.sample().to_csv(original, index=False, sep=";")
            z = Path(temp) / "gol.zip"
            write_zip(z, original, formatted, full, excluded, coupons, final, removed,
                      prior, after, week, stats)
            self.assertTrue(z.exists())
            with zipfile.ZipFile(z) as archive:
                data = archive.read("GOL FINAL.csv").decode("utf-8-sig").splitlines()
                pdf = archive.read("GOL_RELATORIOS.pdf")
            self.assertTrue(pdf.startswith(b"%PDF"))
            self.assertEqual(data[0].split(";"), [
                "DATA_EMISSAO", "DATA_EMBARQUE", "HORA_EMBARQUE", "CIA_BILHETE",
                "TRECHO", "ORIGEM", "DESTINO", "CUPOM", "BILHETE", "TIPO",
                "CLASSE_SERVICO", "CIA_VOO", "VALOR_BRL", "EMPRESA", "CNPJ",
                "IATA_AGENCIA", "BASE_VENDA", "QTD_PAX", "NUM_VOO", "CONSOLIDADA",
            ])
            self.assertIn("100,00", "\n".join(data[1:]))

    def test_advp_scope_and_pivot_totals(self):
        formatted, _, _, _ = format_gol(self.sample(), {57501255})
        groups = {"CURTO": (0, 27), "LONGO": (28, None)}
        records = report_matrix(formatted, {"2026-10-05"}, groups, {"CONS"})
        self.assertEqual(len(records), 10)
        self.assertFalse(any(r["kind"] == "CORP" and r["advp"] != "GERAL" for r in records))
        layout = report_layouts(records, groups, {"CONS"})["SEMANA"]
        self.assertEqual([t["title"] for t in layout], ["RELATÓRIO CONS", "RELATÓRIO CORP", "RELATÓRIO GOV"])
        for t in layout:
            cia_values = [r["cells"][-1]["value"] for r in t["rows"] if r["cia"] != "Total geral"]
            self.assertEqual(cia_values, sorted(cia_values, reverse=True))
        self.assertEqual([b["label"] for b in layout[0]["bands"]],
                         ["ADVP 0 A 27", "ADVP 28+", "ADVP TOTAL"])
        cons_total = next(r for r in records if r["kind"] == "CONS" and r["period"] == "SEMANA" and r["advp"] == "GERAL")
        self.assertAlmostEqual(layout[0]["rows"][-1]["cells"][-1]["value"], cons_total["total"])

    def test_complement_detection_and_period_titles(self):
        formatted, _, _, _ = format_gol(self.sample(), {57501255})
        extra_raw = self.sample().copy()
        extra_raw["DATA_EMISSAO"] = "10/10/2026"
        extra_raw["DATA_EMBARQUE"] = "25/10/2026"
        extra_raw["BILHETE"] = [f"2234567890{i:03}" for i in range(5)]
        extra, status, stats = prepare_complement(extra_raw, {57501255})
        self.assertEqual(status, "FORMATADA AUTOMATICAMENTE")
        self.assertEqual(stats["formatted_rows"], 5)
        already, status, stats = prepare_complement(formatted, {57501255})
        self.assertEqual(status, "JÁ FORMATADA")
        self.assertIsNone(stats)
        full = merge_period(formatted.assign(_SOURCE_PERIOD="ORIGINAL"), extra)
        titles = period_titles(full, {"2026-10-05"}, None, True)
        self.assertEqual(titles["SEMANA"], "RELATÓRIO DA SEMANA DE 05 A 05 OUTUBRO")
        self.assertEqual(titles["COMPLETO"], "RELATÓRIO COMPLETO 05 A 10 OUTUBRO")
        self.assertEqual(list(period_titles(full, {"2026-10-05"}, None, False)), ["SEMANA"])
        matrix = report_matrix(full, {"2026-10-05"}, {}, set(), include_complete=False)
        self.assertFalse(any(r["period"] == "COMPLETO" for r in matrix))

    def test_recover_missing_emission_and_report_base_filters(self):
        source = self.sample()
        missing = source.drop(columns="DATA_EMISSAO").copy()
        missing["DATA_EMBARQUE"] = 46315  # 20/10/2026 no Excel
        recovered = recover_emission_dates(missing, source)
        self.assertEqual(recovered["DATA_EMISSAO"].iloc[0].strftime("%d/%m/%Y"), "05/10/2026")
        self.assertEqual(prepare_complement(recovered, set())[1], "FORMATADA AUTOMATICAMENTE")
        invalid = source.copy()
        invalid.loc[0, "IATA_AGENCIA"] = "99999999"
        with self.assertRaisesRegex(ValueError, "IATA_AGENCIA"):
            recover_emission_dates(missing, invalid)
        formatted = format_gol(source, {57501255})[0]
        options = filter_options(formatted)
        self.assertIn("CGB", options["bases"]["CONS"])
        self.assertIn("CGB-CORP", options["bases"]["CORP"])
        filters = default_filters(formatted)
        filters["bases"]["CONS"] = ["CGB"]
        self.assertTrue(report_rows(formatted, "CONS", filters=filters)["BASE_VENDA"].eq("CGB").all())

    def test_web_flow(self):
        import app as web
        previous = web.JOBS_DIR
        with TemporaryDirectory() as temp:
            web.JOBS_DIR = Path(temp)
            try:
                client = web.app.test_client()
                raw = self.sample().to_csv(index=False, sep=";").encode("utf-8")
                response = client.post("/gol/formatar", data={
                    "arquivo": (io.BytesIO(raw), "gol.csv"),
                    "corp_iatas": "57501255",
                }, content_type="multipart/form-data")
                self.assertEqual(response.status_code, 200)
                match = re.search(rb"/gol/relatorios/([0-9a-f]{32})", response.data)
                self.assertIsNotNone(match)
                job = match.group(1).decode()
                response = client.post(f"/gol/relatorios/{job}", data={
                    "use_advp": "S", "short_start": "0", "short_end": "27",
                    "use_intermediate": "N", "advp_scope": "CONS",
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"ADVP 0 A 27", response.data)
                self.assertIn(b"CORP ADVP GERAL", response.data)
                self.assertIn(b"Total geral", response.data)
                self.assertNotIn(b"RELAT\xc3\x93RIO COMPLETO", response.data)
                response = client.post(f"/gol/filtros/{job}", data={
                    "cupons": "1", "tipos": "N", "bases_CONS": "CGB",
                    "date_start": "2026-10-05", "date_end": "2026-10-05",
                })
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"RELAT\xc3\x93RIO DA SEMANA DE 05 A 05 OUTUBRO", response.data)
                response = client.post(f"/gol/exportar/{job}", data={})
                self.assertEqual(response.status_code, 200)
                self.assertIn(b"ZIP GOL pronto", response.data)
                response = client.get(f"/gol/download/{job}")
                self.assertEqual(response.status_code, 200)
                self.assertTrue(zipfile.is_zipfile(io.BytesIO(response.data)))
                with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                    from openpyxl import load_workbook
                    workbook = load_workbook(io.BytesIO(archive.read("GOL_FORMATADO_COM_DINAMICA.xlsx")), read_only=True)
                    self.assertIn("DINAMICA_SEMANA", workbook.sheetnames)
                    self.assertIn("DIA_A_DIA", workbook.sheetnames)
                    self.assertNotIn("DINAMICA_COMPLETO", workbook.sheetnames)
                    self.assertEqual(workbook["DINAMICA_SEMANA"]["B4"].value, "ADVP 0 A 27")
                    workbook.close()
                response.close()
                complement = self.sample().copy()
                complement["DATA_EMISSAO"] = "10/10/2026"
                complement["DATA_EMBARQUE"] = "25/10/2026"
                complement["BILHETE"] = [f"2234567890{i:03}" for i in range(5)]
                without_emission = complement.drop(columns="DATA_EMISSAO")
                response = client.post(f"/gol/relatorios/{job}", data={
                    "complementar": (io.BytesIO(without_emission.to_csv(index=False, sep=";").encode()), "outro.csv"),
                    "origem_datas": (io.BytesIO(complement.to_csv(index=False, sep=";").encode()), "origem.csv"),
                    "use_advp": "S", "advp_scope": "CONS", "short_start": "0",
                    "short_end": "27", "use_intermediate": "N",
                }, content_type="multipart/form-data")
                self.assertEqual(response.status_code, 200)
                self.assertIn("RELATÓRIO COMPLETO 05 A 10 OUTUBRO".encode(), response.data)
                self.assertIn("FORMATADA AUTOMATICAMENTE".encode(), response.data)
                response = client.post(f"/gol/exportar/{job}", data={})
                self.assertEqual(response.status_code, 200)
                response = client.get(f"/gol/download/{job}")
                with zipfile.ZipFile(io.BytesIO(response.data)) as archive:
                    self.assertIn("01_ORIGINAL/COMPLEMENTAR_outro.csv", archive.namelist())
                    self.assertIn("GOL_RELATORIOS.pdf", archive.namelist())
                    self.assertIn("GOL FINAL.csv", archive.namelist())
                    self.assertIn("GOL PRINCIPAL.csv", archive.namelist())
                    self.assertIn("GOL BASE COMPLETA.csv", archive.namelist())
                    self.assertIn("GOL_COMPARATIVO_DIA_A_DIA.csv", archive.namelist())
                    from openpyxl import load_workbook
                    workbook = load_workbook(io.BytesIO(archive.read("GOL_FORMATADO_COM_DINAMICA.xlsx")), read_only=True)
                    self.assertIn("DINAMICA_COMPLETO", workbook.sheetnames)
                    self.assertIn("DIA_A_DIA", workbook.sheetnames)
                    workbook.close()
                response.close()
            finally:
                web.JOBS_DIR = previous


    def test_daily_comparison_tables_structure(self):
        df = pd.DataFrame({
            "DATA_EMISSAO": ["05/10/2026", "05/10/2026", "06/10/2026", "06/10/2026"],
            "DATA_EMBARQUE": ["10/10/2026", "10/11/2026", "10/10/2026", "10/11/2026"],
            "CIA_BILHETE": ["JJ", "JJ", "G3", "JJ"],
            "CUPOM": ["1"] * 4,
            "BILHETE": [f"1234567890{i:03}" for i in range(4)],
            "TIPO": ["N"] * 4,
            "ORIGEM": ["CGB"] * 4,
            "DESTINO": ["GRU"] * 4,
            "TRECHO": ["CGB/GRU"] * 4,
            "IATA_AGENCIA": ["57500001"] * 4,
            "VALOR_BRL": ["200,00", "200,00", "150,00", "200,00"],
            "BASE_VENDA": ["CGB"] * 4,
            "NUM_VOO": ["1234"] * 4,
            "CABINE": ["Y"] * 4,
            "CONSOLIDADA": ["1"] * 4,
            "CONSOLIDADA.1": ["ANTIGA"] * 4,
            "GOV": ["Nao"] * 4,
        })
        formatted, _, _, _ = format_gol(df, set())
        groups = {"CURTO": (0, 27), "LONGO": (28, None)}
        week = {"2026-10-05", "2026-10-06"}
        tables = daily_comparison_tables(formatted, formatted, week, groups, {"CONS"})
        titles = [t["title"] for t in tables]
        self.assertIn("CONS — ADVP 0 A 27 POR DIA", titles)
        self.assertIn("CONS — ADVP 28+ POR DIA", titles)
        self.assertIn("CONS — ADVP TOTAL POR DIA", titles)
        self.assertIn("CORP — ADVP TOTAL POR DIA", titles)
        self.assertIn("GOV — ADVP TOTAL POR DIA", titles)
        cons_curto = next(t for t in tables if "0 A 27" in t["title"])
        self.assertEqual(cons_curto["rows"][-1]["date_label"], "TOTAL DA SEMANA")
        self.assertTrue(cons_curto["rows"][-1]["is_total"])
        csv_df = daily_comparison_dataframe(tables)
        self.assertIn("JJ_SHARE_ORIGINAL", csv_df.columns)
        self.assertIn("TOTAL_ORIGINAL_BRL", csv_df.columns)

    def test_longo_and_daily_latam_guard(self):
        # Gera vendas em que JJ tem bastante folga inicialmente no longo e curto
        rows = []
        # Curto: 2 vendas JJ (R$ 500 cada), 1 venda G3 (R$ 300) -> total 1300
        # Longo: 4 vendas JJ (R$ 500 cada), 1 venda G3 (R$ 300) -> JJ = 2000, G3 = 300
        for i in range(2):
            rows.append({"DATA_EMISSAO": "05/10/2026", "DATA_EMBARQUE": "10/10/2026", "CIA_BILHETE": "JJ",
                         "CUPOM": "1", "BILHETE": f"100000000{i:04}", "TIPO": "N", "ORIGEM": "CGB",
                         "DESTINO": "GRU", "TRECHO": "CGB/GRU", "IATA_AGENCIA": "57500001",
                         "VALOR_BRL": "500,00", "BASE_VENDA": "CGB", "NUM_VOO": "100", "CABINE": "Y",
                         "CONSOLIDADA": "1", "CONSOLIDADA.1": "A", "GOV": "Nao"})
        rows.append({"DATA_EMISSAO": "05/10/2026", "DATA_EMBARQUE": "10/10/2026", "CIA_BILHETE": "G3",
                     "CUPOM": "1", "BILHETE": "100000000099", "TIPO": "N", "ORIGEM": "CGB",
                     "DESTINO": "GRU", "TRECHO": "CGB/GRU", "IATA_AGENCIA": "57500001",
                     "VALOR_BRL": "300,00", "BASE_VENDA": "CGB", "NUM_VOO": "100", "CABINE": "Y",
                     "CONSOLIDADA": "1", "CONSOLIDADA.1": "A", "GOV": "Nao"})
        for i in range(4):
            rows.append({"DATA_EMISSAO": "06/10/2026", "DATA_EMBARQUE": "20/11/2026", "CIA_BILHETE": "JJ",
                         "CUPOM": "1", "BILHETE": f"200000000{i:04}", "TIPO": "N", "ORIGEM": "CGB",
                         "DESTINO": "GRU", "TRECHO": "CGB/GRU", "IATA_AGENCIA": "57500001",
                         "VALOR_BRL": "500,00", "BASE_VENDA": "CGB", "NUM_VOO": "100", "CABINE": "Y",
                         "CONSOLIDADA": "1", "CONSOLIDADA.1": "A", "GOV": "Nao"})
        rows.append({"DATA_EMISSAO": "06/10/2026", "DATA_EMBARQUE": "20/11/2026", "CIA_BILHETE": "G3",
                     "CUPOM": "1", "BILHETE": "200000000099", "TIPO": "N", "ORIGEM": "CGB",
                     "DESTINO": "GRU", "TRECHO": "CGB/GRU", "IATA_AGENCIA": "57500001",
                     "VALOR_BRL": "300,00", "BASE_VENDA": "CGB", "NUM_VOO": "100", "CABINE": "Y",
                     "CONSOLIDADA": "1", "CONSOLIDADA.1": "A", "GOV": "Nao"})

        df = pd.DataFrame(rows)
        formatted, _, _, _ = format_gol(df, set())
        groups = {"CURTO": (0, 27), "LONGO": (28, None)}
        week = {"2026-10-05", "2026-10-06"}
        # Pedir meta alta para G3 (ex: 45%), sem meta para Longo
        final, removed, _, after = adjust_shares(
            formatted, week, {"CONS": {"general": 0.45, "week": None, "advp": {}}},
            set(), groups, advp_kinds={"CONS"}
        )
        # Verificar que no LONGO, LATAM (JJ) continua na frente da GOL (G3) com pelo menos 3 pp
        longo_after = next(r for r in after if r["kind"] == "CONS" and r["advp"] == "LONGO" and r["period"] == "SEMANA")
        self.assertGreaterEqual(longo_after["shares"]["JJ"], longo_after["shares"]["G3"] + 0.03 - 1e-7)
        # Verificar que no CURTO, LATAM também continua na frente
        curto_after = next(r for r in after if r["kind"] == "CONS" and r["advp"] == "CURTO" and r["period"] == "SEMANA")
        self.assertGreaterEqual(curto_after["shares"]["JJ"], curto_after["shares"]["G3"] + 0.03 - 1e-7)


if __name__ == "__main__":
    unittest.main()
