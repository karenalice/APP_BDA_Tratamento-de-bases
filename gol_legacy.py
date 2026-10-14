import csv
import os
import random
import re
import string
import unicodedata
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURAÇÕES
# ============================================================
VERSAO_PROGRAMA = "4.0 - BILHETE, CONSOLIDADA E VALOR_BRL"
IATAS_EXCLUIR = {"57516197", "57516196"}
IATAS_CORP = {
    "57545666",
    "57560753",
    "57593233",
    "57562805",
    "57593244",
    "57545935",
    "57545681",
    "57559972",
    "57556844",
    "57562352",
    "57593266",
    "57562293",
    "57527035",
    "57559950",
    "57562341",
    "57593281",
}
PREFIXOS_BILHETE_EXCLUIR = ("12744", "12745", "95744", "95745")
CNPJ_PADRAO = "10000000000000"

COLUNAS_FINAIS = [
    "DATA_EMISSAO",
    "DATA_EMBARQUE",
    "HORA_EMBARQUE",
    "CIA_BILHETE",
    "TRECHO",
    "ORIGEM",
    "DESTINO",
    "CUPOM",
    "BILHETE",
    "TIPO",
    "CLASSE_SERVICO",
    "CIA_VOO",
    "VALOR_BRL",
    "EMPRESA",
    "CNPJ",
    "IATA_AGENCIA",
    "BASE_VENDA",
    "QTD_PAX",
    "NUM_VOO",
    "CONSOLIDADA",
]

# CONSOLIDADA não entra nesta lista porque deve permanecer no resultado final.
COLUNAS_EXCLUIR = [
    "DATA_EXTRACAO",
    "HORA_EMISSAO",
    "DATA_RESERVA",
    "HORA_RESERVA",
    "HORA_POUSO",
    "BASE_TARIFARIA",
    "TKT_DESIG",
    "FAMILIA_TARIFARIA",
    "CLASSE_TARIFA",
    "CLASSE_SERVICO",  # A coluna antiga é ignorada; CABINE vira CLASSE_SERVICO.
    "OND_DIRECIONAL",
    "TOUR_CODE",
    "RT_OW",
    "VALOR_USD",
    "TARIFA_PUBLICA_BRL",
    "TARIFA_PUBLICA_USD",
    "PNR_AGENCIA",
    "PNR_CIA_AEREA",
    "SELFBOOKING_OFFLINE",
    "NOME_PAX",
    "TIPO_PAX",
    "CPF_PAX",
    "E_MAIL_PAX",
    "CELULAR_PAX",
    "TIER_FIDELIDADE_PAX",
    "TIPO_PAGAMENTO",
    "DIGITOS_CC",
    "GRUPO_EMPRESA",
    "GRUPO_CONSOLIDADA",
]

COLUNAS_OBRIGATORIAS = [
    "DATA_EMBARQUE",
    "DATA_EMISSAO",
    "CUPOM",
    "ORIGEM",
    "DESTINO",
    "TRECHO",
    "IATA_AGENCIA",
    "VALOR_BRL",
    "BILHETE",
    "BASE_VENDA",
    "NUM_VOO",
]

GERADOR_ALEATORIO = random.SystemRandom()
CARACTERES_INVALIDOS_EXCEL = re.compile(r"[\x00-\x08\x0B\x0C\x0E-\x1F]+")


# ============================================================
# FUNÇÕES AUXILIARES
# ============================================================
def normalizar_nome_coluna(nome):
    """Remove acentos e padroniza o título como MAIUSCULO_COM_UNDERLINE."""
    texto = str(nome).strip()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(r"[^A-Za-z0-9]+", "_", texto).strip("_")
    return texto.upper()


def tornar_nomes_unicos(nomes):
    """Evita erro quando o arquivo possui duas colunas com o mesmo título."""
    contagem = {}
    resultado = []

    for nome in nomes:
        quantidade = contagem.get(nome, 0)
        resultado.append(nome if quantidade == 0 else f"{nome}_{quantidade}")
        contagem[nome] = quantidade + 1

    return resultado


def normalizar_texto(valor):
    if pd.isna(valor):
        return ""

    texto = str(valor).strip()
    if texto.lower() in {"nan", "none", "null", "nat"}:
        return ""

    return texto


def limpar_caracteres_invalidos_excel(valor):
    """Remove controles invisiveis que o Excel/OpenXML nao aceita em celulas."""
    if not isinstance(valor, str):
        return valor

    return CARACTERES_INVALIDOS_EXCEL.sub(" ", valor)


def preparar_dataframe_para_excel(df):
    """Limpa textos antes da exportacao sem alterar numeros, datas ou nulos."""
    resultado = df.copy()
    colunas_textuais = resultado.select_dtypes(include=["object", "string"]).columns

    for coluna in colunas_textuais:
        resultado[coluna] = resultado[coluna].map(limpar_caracteres_invalidos_excel)

    return resultado


def remover_acentos_maiusculo(valor):
    texto = normalizar_texto(valor)
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto.upper()


def esta_vazio(valor):
    return normalizar_texto(valor) == ""


def somente_digitos(valor):
    """Mantém somente dígitos sem acrescentar o sufixo .0."""
    if pd.isna(valor):
        return ""

    if isinstance(valor, bool):
        return "1" if valor else "0"

    if isinstance(valor, int):
        return str(valor)

    if isinstance(valor, float):
        if pd.isna(valor):
            return ""
        if valor.is_integer():
            return str(int(valor))

    texto = str(valor).strip()
    texto = re.sub(r"\.0+$", "", texto)
    return re.sub(r"\D", "", texto)


def converter_numero_brasileiro(valor):
    """Converte 1.234,56, 1234.56 ou número simples para float."""
    if pd.isna(valor) or valor == "":
        return float("nan")

    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        return float(valor)

    texto = str(valor).strip()
    texto = texto.replace("R$", "").replace(" ", "")

    if not texto:
        return float("nan")

    if "," in texto and "." in texto:
        if texto.rfind(",") > texto.rfind("."):
            texto = texto.replace(".", "").replace(",", ".")
        else:
            texto = texto.replace(",", "")
    elif "," in texto:
        texto = texto.replace(".", "").replace(",", ".")

    texto = re.sub(r"[^0-9.\-]", "", texto)

    try:
        return float(texto)
    except ValueError:
        return float("nan")


def detectar_separador_csv(caminho, encoding):
    with open(caminho, "r", encoding=encoding, errors="strict", newline="") as arquivo:
        amostra = arquivo.read(8192)

    try:
        return csv.Sniffer().sniff(amostra, delimiters=";,\t|").delimiter
    except csv.Error:
        # Os arquivos desta origem normalmente utilizam ponto e vírgula.
        return ";"


def ler_arquivo(caminho):
    extensao = Path(caminho).suffix.lower()

    if extensao == ".xlsx":
        return pd.read_excel(caminho, engine="openpyxl", dtype=object)

    if extensao == ".xls":
        return pd.read_excel(caminho, engine="xlrd", dtype=object)

    if extensao == ".csv":
        erros = []

        for encoding in ("utf-8-sig", "utf-8", "latin-1", "cp1252"):
            try:
                separador = detectar_separador_csv(caminho, encoding)
                return pd.read_csv(
                    caminho,
                    sep=separador,
                    encoding=encoding,
                    dtype=object,
                    low_memory=False,
                )
            except Exception as erro:
                erros.append(f"{encoding}: {erro}")

        raise ValueError(
            "Não foi possível ler o CSV. Tentativas realizadas:\n"
            + "\n".join(erros)
        )

    raise ValueError("Formato não aceito. Selecione um arquivo XLS, XLSX ou CSV.")


def primeira_coluna_preenchida(df, colunas):
    """Combina colunas da esquerda para a direita, usando o primeiro valor preenchido."""
    if not colunas:
        return pd.Series("", index=df.index, dtype=object)

    resultado = pd.Series("", index=df.index, dtype=object)

    for coluna in colunas:
        valores = df[coluna].apply(normalizar_texto)
        mascara = resultado.eq("") & valores.ne("")
        resultado.loc[mascara] = valores.loc[mascara]

    return resultado


def preparar_colunas(df):
    """
    Padroniza nomes e resolve as variações encontradas no arquivo:
    - IATA ou IATA_AGENCIA passam a ser IATA_AGENCIA.
    - CONSOLIDADA duplicada não interrompe o programa.
    """
    nomes_normalizados = [normalizar_nome_coluna(coluna) for coluna in df.columns]
    df.columns = tornar_nomes_unicos(nomes_normalizados)

    # Reconhece variações do nome da coluna de IATA da agência.
    # A ordem é importante: prioriza IATA_AGENCIA e só depois IATA.
    aliases_iata = {
        "IATA_AGENCIA",
        "IATAAGENCIA",
        "IATA_DA_AGENCIA",
        "NUMERO_IATA_AGENCIA",
        "NUM_IATA_AGENCIA",
        "NUMERO_IATA",
        "NUM_IATA",
        "IATA",
    }

    colunas_iata = []

    # Primeiro procura os títulos mais específicos.
    for coluna in df.columns:
        base_coluna = re.sub(r"_\d+$", "", coluna)
        if base_coluna in aliases_iata:
            colunas_iata.append(coluna)

    # Como último recurso, aceita uma coluna que contenha IATA e AGENCIA.
    if not colunas_iata:
        colunas_iata = [
            coluna
            for coluna in df.columns
            if "IATA" in coluna and "AGENCIA" in coluna
        ]

    if colunas_iata:
        # Ordena para IATA_AGENCIA ter prioridade sobre aliases genéricos.
        colunas_iata.sort(
            key=lambda coluna: (
                0 if coluna == "IATA_AGENCIA" else
                1 if coluna.startswith("IATA_AGENCIA_") else
                2 if "AGENCIA" in coluna else
                3
            )
        )
        df["IATA_AGENCIA"] = primeira_coluna_preenchida(df, colunas_iata)

    # Quando existem duas colunas CONSOLIDADA, a correta é a SEGUNDA.
    # Após tornar os nomes únicos, elas ficam como CONSOLIDADA e CONSOLIDADA_1.
    colunas_consolidada = [
        coluna
        for coluna in df.columns
        if coluna == "CONSOLIDADA"
        or re.fullmatch(r"CONSOLIDADA_\d+", coluna)
    ]

    if colunas_consolidada:
        fonte_consolidada = (
            colunas_consolidada[1]
            if len(colunas_consolidada) >= 2
            else colunas_consolidada[0]
        )
        df["CONSOLIDADA"] = df[fonte_consolidada].apply(normalizar_texto)

    return df


def validar_colunas(df):
    faltantes = [
        coluna for coluna in COLUNAS_OBRIGATORIAS if coluna not in df.columns
    ]

    # CABINE é a fonte preferencial. Se não existir, aceita CLASSE_SERVICO.
    if "CABINE" not in df.columns and "CLASSE_SERVICO" not in df.columns:
        faltantes.append("CABINE ou CLASSE_SERVICO")

    if faltantes:
        disponiveis = ", ".join(df.columns)
        raise ValueError(
            "As seguintes colunas obrigatórias não foram encontradas:\n\n"
            + "\n".join(f"- {coluna}" for coluna in faltantes)
            + "\n\nPara a agência, são aceitos títulos como "
              "IATA_AGENCIA, IATA, NUMERO_IATA ou NUM_IATA."
            + "\n\nColunas encontradas no arquivo:\n"
            + disponiveis
        )


def gerar_numero_voo(valor):
    """
    Retorna um número de voo com exatamente quatro dígitos.
    - 4 dígitos válidos: mantém.
    - Mais de 4: usa os quatro primeiros.
    - Menos de 4, vazio ou inválido: gera número fictício entre 1000 e 9999.
    """
    digitos = somente_digitos(valor)

    if len(digitos) >= 4:
        candidato = digitos[:4]
        if candidato[0] != "0":
            return int(candidato)

    return GERADOR_ALEATORIO.randint(1000, 9999)


def gerar_classe_servico(valor):
    texto = normalizar_texto(valor)
    if texto:
        return texto
    return GERADOR_ALEATORIO.choice(string.ascii_uppercase)


def tratar_bilhete(valor):
    """Retorna o bilhete como texto com 13 dígitos, sem apóstrofo visível."""
    digitos = somente_digitos(valor)
    return digitos.zfill(13)[:13]


def tratar_trecho(valor):
    texto = normalizar_texto(valor)
    texto = re.sub(r"\s+", "", texto)
    texto = re.sub(r"/{2,}", "/", texto)
    return texto


def montar_motivos_exclusao(df):
    """Retorna todos os motivos que provocaram a exclusão de cada linha."""
    motivos = pd.Series(
        [[] for _ in range(len(df))],
        index=df.index,
        dtype=object,
    )

    data_emissao = pd.to_datetime(
        df["DATA_EMISSAO"],
        errors="coerce",
        dayfirst=True,
    )
    data_embarque = pd.to_datetime(
        df["DATA_EMBARQUE"],
        errors="coerce",
        dayfirst=True,
    )
    cupom_numerico = pd.to_numeric(df["CUPOM"], errors="coerce")
    valor_brl_numerico = df["VALOR_BRL"].apply(converter_numero_brasileiro)
    iata_digitos = df["IATA_AGENCIA"].apply(somente_digitos)
    bilhete_digitos = df["BILHETE"].apply(somente_digitos)

    regras = [
        (
            data_embarque.notna()
            & data_emissao.notna()
            & data_embarque.lt(data_emissao),
            "DATA_EMBARQUE menor que DATA_EMISSAO",
        ),
        (cupom_numerico.eq(0), "CUPOM igual a 0"),
        (df["ORIGEM"].apply(esta_vazio), "ORIGEM vazia"),
        (df["DESTINO"].apply(esta_vazio), "DESTINO vazio"),
        (df["TRECHO"].apply(esta_vazio), "TRECHO vazio"),
        (iata_digitos.isin(IATAS_EXCLUIR), "IATA_AGENCIA bloqueado"),
        (
            valor_brl_numerico.notna()
            & valor_brl_numerico.abs().le(0.000000001),
            "VALOR_BRL igual a 0",
        ),
        (
            bilhete_digitos.str.startswith(PREFIXOS_BILHETE_EXCLUIR, na=False),
            "Prefixo de BILHETE bloqueado",
        ),
    ]

    for mascara, descricao in regras:
        for indice in df.index[mascara.fillna(False)]:
            motivos.at[indice].append(descricao)

    return motivos


def corrigir_cupons(df):
    """
    Renumera todos os cupons por BILHETE na ordem original das linhas.
    Cada bilhete passa a ter a sequência 1, 2, 3, 4...
    """
    cupom_anterior = df["CUPOM"].copy()

    df["CUPOM"] = (
        df.groupby("BILHETE", sort=False, dropna=False)
        .cumcount()
        .add(1)
        .astype("Int64")
    )

    anterior_texto = cupom_anterior.apply(normalizar_texto)
    novo_texto = df["CUPOM"].astype("string").fillna("")

    alterados = anterior_texto.ne(novo_texto)

    auditoria = pd.DataFrame(
        {
            "LINHA_ORIGINAL": df.loc[alterados, "_LINHA_ORIGINAL"],
            "BILHETE": df.loc[alterados, "BILHETE"],
            "CUPOM_ANTERIOR": cupom_anterior.loc[alterados],
            "CUPOM_CORRIGIDO": df.loc[alterados, "CUPOM"],
        }
    )

    return df, auditoria


def ajustar_largura_e_estilo(writer, nome_aba, df):
    ws = writer.book[nome_aba]
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    preenchimento = PatternFill(fill_type="solid", fgColor="D9EAF7")
    fonte = Font(bold=True)

    for celula in ws[1]:
        celula.fill = preenchimento
        celula.font = fonte

    for indice, coluna in enumerate(df.columns, start=1):
        valores = [str(coluna)]
        valores.extend(str(v) for v in df[coluna].head(500).fillna(""))
        largura = min(max(len(v) for v in valores) + 2, 45)
        ws.column_dimensions[get_column_letter(indice)].width = max(largura, 10)

    for nome_coluna in ("BILHETE", "CNPJ", "IATA_AGENCIA"):
        if nome_coluna in df.columns:
            numero_coluna = df.columns.get_loc(nome_coluna) + 1
            for linha in range(2, len(df) + 2):
                ws.cell(linha, numero_coluna).number_format = "@"


# ============================================================
# PROCESSAMENTO PRINCIPAL
# ============================================================
def processar_arquivo(caminho_entrada, caminho_saida):
    df = ler_arquivo(caminho_entrada)
    df = preparar_colunas(df)
    validar_colunas(df)

    df.insert(0, "_LINHA_ORIGINAL", range(2, len(df) + 2))
    total_inicial = len(df)

    # ------------------------
    # EXCLUSÕES
    # ------------------------
    motivos = montar_motivos_exclusao(df)
    mascara_excluir = motivos.apply(bool)

    linhas_excluidas = df.loc[mascara_excluir].copy()
    linhas_excluidas.insert(
        0,
        "MOTIVO_EXCLUSAO",
        motivos.loc[mascara_excluir].apply(lambda lista: "; ".join(lista)),
    )

    df = df.loc[~mascara_excluir].copy()

    # ------------------------
    # TRATAMENTOS
    # ------------------------
    df["DATA_EMISSAO"] = pd.to_datetime(
        df["DATA_EMISSAO"], errors="coerce", dayfirst=True
    )
    df["DATA_EMBARQUE"] = pd.to_datetime(
        df["DATA_EMBARQUE"], errors="coerce", dayfirst=True
    )

    df["VALOR_BRL"] = df["VALOR_BRL"].apply(converter_numero_brasileiro)

    # BASE_VENDA sem espaços; GOV tem prioridade sobre o sufixo CORP.
    base_sem_espacos = (
        df["BASE_VENDA"]
        .apply(normalizar_texto)
        .str.replace(r"\s+", "", regex=True)
    )
    iata_corp = df["IATA_AGENCIA"].apply(somente_digitos).isin(IATAS_CORP)
    if "GOV" in df.columns:
        gov_sim = df["GOV"].apply(remover_acentos_maiusculo).eq("SIM")
    else:
        gov_sim = pd.Series(False, index=df.index)

    termina_corp = base_sem_espacos.str.upper().str.endswith("-CORP")
    base_sem_espacos = base_sem_espacos.where(
        ~iata_corp | ~termina_corp,
        base_sem_espacos.str.replace(r"-CORP$", "", case=False, regex=True),
    )

    precisa_sufixo = gov_sim & ~base_sem_espacos.str.upper().str.endswith("-GOV")
    df["BASE_VENDA"] = base_sem_espacos.where(
        ~precisa_sufixo,
        base_sem_espacos + "-GOV",
    )
    precisa_corp = (
        iata_corp
        & ~gov_sim
        & ~df["BASE_VENDA"].str.upper().str.endswith("-CORP")
    )
    df["BASE_VENDA"] = df["BASE_VENDA"].where(
        ~precisa_corp,
        df["BASE_VENDA"] + "-CORP",
    )

    # CNPJ fixo com 14 caracteres.
    df["CNPJ"] = CNPJ_PADRAO

    # NUM_VOO com quatro dígitos.
    df["NUM_VOO"] = df["NUM_VOO"].apply(gerar_numero_voo).astype(int)

    # CABINE vira CLASSE_SERVICO. Caso CABINE não exista, usa a antiga.
    if "CABINE" in df.columns:
        fonte_classe = df["CABINE"]
    else:
        fonte_classe = df["CLASSE_SERVICO"]

    df["CLASSE_SERVICO"] = fonte_classe.apply(gerar_classe_servico)

    # Bilhete como texto com 13 dígitos, sem mostrar apóstrofo na célula.
    df["BILHETE"] = df["BILHETE"].apply(tratar_bilhete).astype("string")

    # Trecho sem espaços e sem //.
    df["TRECHO"] = df["TRECHO"].apply(tratar_trecho)

    # IATA no título correto e como texto.
    df["IATA_AGENCIA"] = df["IATA_AGENCIA"].apply(somente_digitos)

    # Cupons corrigidos automaticamente por bilhete.
    df, cupons_corrigidos = corrigir_cupons(df)

    # ------------------------
    # COLUNAS FINAIS
    # ------------------------
    for coluna in COLUNAS_FINAIS:
        if coluna not in df.columns:
            df[coluna] = ""

    base_final = df[COLUNAS_FINAIS].copy()

    # Conferência final: nenhuma linha com VALOR_BRL zero pode permanecer.
    valor_final = base_final["VALOR_BRL"].apply(converter_numero_brasileiro)
    mascara_zero_final = valor_final.notna() & valor_final.abs().le(0.000000001)
    if mascara_zero_final.any():
        raise RuntimeError(
            "Falha de validação: ainda existem linhas com VALOR_BRL igual a zero."
        )

    # Resumo das exclusões.
    contagem_motivos = {}
    for lista in motivos:
        for motivo in lista:
            contagem_motivos[motivo] = contagem_motivos.get(motivo, 0) + 1

    resumo = pd.DataFrame(
        [
            {"INDICADOR": "Linhas no arquivo original", "QUANTIDADE": total_inicial},
            {
                "INDICADOR": "Linhas excluídas (total único)",
                "QUANTIDADE": int(mascara_excluir.sum()),
            },
            {"INDICADOR": "Linhas na base tratada", "QUANTIDADE": len(base_final)},
            {
                "INDICADOR": "Cupons alterados",
                "QUANTIDADE": len(cupons_corrigidos),
            },
        ]
        + [
            {"INDICADOR": f"Exclusão: {motivo}", "QUANTIDADE": quantidade}
            for motivo, quantidade in contagem_motivos.items()
        ]
    )

    # Remove caracteres de controle invisiveis rejeitados pelo openpyxl.
    base_final = preparar_dataframe_para_excel(base_final)
    linhas_excluidas = preparar_dataframe_para_excel(linhas_excluidas)
    cupons_corrigidos = preparar_dataframe_para_excel(cupons_corrigidos)
    resumo = preparar_dataframe_para_excel(resumo)

    # ------------------------
    # EXPORTAÇÃO
    # ------------------------
    with pd.ExcelWriter(
        caminho_saida,
        engine="openpyxl",
        date_format="DD/MM/YYYY",
        datetime_format="DD/MM/YYYY",
    ) as writer:
        base_final.to_excel(writer, sheet_name="BASE_TRATADA", index=False)
        linhas_excluidas.to_excel(
            writer, sheet_name="LINHAS_EXCLUIDAS", index=False
        )
        cupons_corrigidos.to_excel(
            writer, sheet_name="CUPONS_CORRIGIDOS", index=False
        )
        resumo.to_excel(writer, sheet_name="RESUMO", index=False)

        ajustar_largura_e_estilo(writer, "BASE_TRATADA", base_final)
        ajustar_largura_e_estilo(
            writer, "LINHAS_EXCLUIDAS", linhas_excluidas
        )
        ajustar_largura_e_estilo(
            writer, "CUPONS_CORRIGIDOS", cupons_corrigidos
        )
        ajustar_largura_e_estilo(writer, "RESUMO", resumo)

    return {
        "total_inicial": total_inicial,
        "total_excluido": int(mascara_excluir.sum()),
        "total_final": len(base_final),
        "cupons_corrigidos": len(cupons_corrigidos),
        "arquivo_saida": caminho_saida,
    }


# ============================================================
# INTERFACE GRÁFICA
# ============================================================
def selecionar_e_processar():
    caminho_entrada = filedialog.askopenfilename(
        title="Selecione a tabela",
        filetypes=[
            ("Tabelas", "*.xls *.xlsx *.csv"),
            ("Excel antigo", "*.xls"),
            ("Excel", "*.xlsx"),
            ("CSV", "*.csv"),
            ("Todos os arquivos", "*.*"),
        ],
    )

    if not caminho_entrada:
        return

    sugestao = f"{Path(caminho_entrada).stem}_TRATADO.xlsx"
    caminho_saida = filedialog.asksaveasfilename(
        title="Salvar tabela tratada",
        initialdir=str(Path(caminho_entrada).parent),
        initialfile=sugestao,
        defaultextension=".xlsx",
        filetypes=[("Arquivo Excel", "*.xlsx")],
    )

    if not caminho_saida:
        return

    try:
        resultado = processar_arquivo(caminho_entrada, caminho_saida)

        messagebox.showinfo(
            "Processamento concluído",
            "Arquivo processado com sucesso!\n\n"
            f"Linhas originais: {resultado['total_inicial']:,}\n"
            f"Linhas excluídas: {resultado['total_excluido']:,}\n"
            f"Linhas finais: {resultado['total_final']:,}\n"
            f"Cupons corrigidos: {resultado['cupons_corrigidos']:,}\n\n"
            f"Salvo em:\n{resultado['arquivo_saida']}",
        )

    except Exception as erro:
        messagebox.showerror(
            "Erro",
            f"Não foi possível processar o arquivo:\n\n{erro}",
        )


def criar_interface():
    janela = tk.Tk()
    janela.title(f"Tratador de Tabelas - versão {VERSAO_PROGRAMA}")
    janela.geometry("600x300")
    janela.resizable(False, False)

    titulo = tk.Label(
        janela,
        text=f"Tratamento de XLS, XLSX e CSV\nVersão {VERSAO_PROGRAMA}",
        font=("Arial", 16, "bold"),
    )
    titulo.pack(pady=(28, 10))

    descricao = tk.Label(
        janela,
        text=(
            "Selecione a tabela para aplicar as exclusões, corrigir os cupons\n"
            "e gerar o arquivo Excel com as colunas finais."
        ),
        font=("Arial", 10),
        justify="center",
    )
    descricao.pack(pady=8)

    botao = tk.Button(
        janela,
        text="Selecionar tabela e processar",
        command=selecionar_e_processar,
        font=("Arial", 11, "bold"),
        width=32,
        height=2,
        cursor="hand2",
    )
    botao.pack(pady=18)

    janela.mainloop()


if __name__ == "__main__":
    criar_interface()
